import os
from docx import Document
from app.parser import QuestionParser
from app.exporter import CBTExporter
from app.batch_engine import BatchEngine

def create_sample_docx(filepath: str):
    """Creates a sample messy Word document to test the converter."""
    doc = Document()
    doc.add_paragraph("Sample Teacher Draft - Physics & General Science")
    
    # Q1: Standard with Ans tag
    doc.add_paragraph("1. What is the capital city of Nigeria?")
    doc.add_paragraph("A) Lagos")
    doc.add_paragraph("B) Abuja")
    doc.add_paragraph("C) Ibadan")
    doc.add_paragraph("D) Kano")
    doc.add_paragraph("ANSWER: B")

    # Q2: Bold correct option without explicit ANSWER tag
    p2 = doc.add_paragraph("2. Which organ in the human body is responsible for pumping blood?")
    doc.add_paragraph("A) Lungs")
    doc.add_paragraph("B) Liver")
    p_c = doc.add_paragraph()
    r = p_c.add_run("C) Heart")
    r.bold = True # Bolded answer
    doc.add_paragraph("D) Kidney")

    # Q3: Inline options on single line
    doc.add_paragraph("3. What is the chemical formula for water?")
    doc.add_paragraph("A) CO2  B) NaCl  C) H2O  D) O2")
    doc.add_paragraph("Ans: C")

    # Q4: With asterisk on answer
    doc.add_paragraph("4. What is the value of 15 multiplied by 4?")
    doc.add_paragraph("A) 45")
    doc.add_paragraph("B) 50")
    doc.add_paragraph("C) 55")
    doc.add_paragraph("*D) 60") # Starred option

    # Q5: Period style numbering & options
    doc.add_paragraph("5. Which continent is Nigeria located in?")
    doc.add_paragraph("A. Asia")
    doc.add_paragraph("B. Africa")
    doc.add_paragraph("C. Europe")
    doc.add_paragraph("D. South America")
    doc.add_paragraph("Correct Answer: B")

    doc.save(filepath)
    print(f"[OK] Created test sample: {filepath}")

def test_full_pipeline():
    test_docx = "test_files/sample_raw_exam.docx"
    create_sample_docx(test_docx)

    # 1. Test Parser
    with open(test_docx, "rb") as f:
        data = f.read()

    parser = QuestionParser()
    result = parser.parse_docx(data)
    questions = result["questions"]
    print(f"\n[Test Result] Total questions detected: {len(questions)}")
    assert len(questions) == 5, f"Expected 5 questions, got {len(questions)}"

    for idx, q in enumerate(questions, 1):
        print(f"  Q{idx}: {q['stem'][:45]}... -> Options: {list(q['options'].keys())} -> Ans: {q['answer']}")
        assert len(q["options"]) == 4, f"Q{idx} should have 4 options"
        assert q["answer"] is not None, f"Q{idx} should have detected an answer"

    # 2. Test Exporter
    exporter = CBTExporter()
    docx_out = exporter.generate_docx(
        questions=questions,
        school_name="HONEYGOLD SCHOOL",
        template_title="CBT EXAM QUESTIONS TEMPLATE"
    )
    assert len(docx_out) > 1000, "Docx generation failed"
    out_docx_path = "test_files/HoneyGold_Formatted_Exam.docx"
    with open(out_docx_path, "wb") as f:
        f.write(docx_out)
    print(f"[OK] Generated HoneyGold CBT Word file: {out_docx_path}")

    aiken_out = exporter.generate_aiken_txt(questions, school_name="HONEYGOLD SCHOOL")
    assert "ANSWER: B" in aiken_out, "Aiken export failed"
    out_aiken_path = "test_files/HoneyGold_Exam_Aiken.txt"
    with open(out_aiken_path, "w", encoding="utf-8") as f:
        f.write(aiken_out)
    print(f"[OK] Generated Aiken CBT text file: {out_aiken_path}")

    excel_out = exporter.generate_excel(questions, school_name="HONEYGOLD SCHOOL")
    assert len(excel_out) > 1000, "Excel generation failed"
    out_excel_path = "test_files/HoneyGold_Exam_Matrix.xlsx"
    with open(out_excel_path, "wb") as f:
        f.write(excel_out)
    print(f"[OK] Generated CBT Excel Matrix file: {out_excel_path}")

    # 3. Test Batch Engine
    batch = BatchEngine()
    summary, zip_bytes = batch.process_multiple_files(
        files_data=[
            ("Science_SS1.docx", data),
            ("Science_SS2.docx", data)
        ],
        school_name="HONEYGOLD SCHOOL"
    )
    assert len(summary) == 2
    assert len(zip_bytes) > 2000
    out_zip_path = "test_files/Batch_Test_Exams.zip"
    with open(out_zip_path, "wb") as f:
        f.write(zip_bytes)
    print(f"[OK] Generated Batch ZIP archive: {out_zip_path}")

    print("\n>>> ALL TESTS PASSED SUCCESSFULLY 100%! <<<")

if __name__ == "__main__":
    test_full_pipeline()
